"""Best-effort text extraction for indexing/tagging/near-duplicate detection."""
from pathlib import Path


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    try:
        if ext in (".txt", ".md", ".tex"):
            return path.read_text(encoding="utf-8", errors="ignore")
        if ext == ".pdf":
            return _extract_pdf(path)
        if ext == ".docx":
            return _extract_docx(path)
        if ext == ".pptx":
            return _extract_pptx(path)
    except Exception:
        return ""
    return ""


def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        return ""
    reader = PdfReader(str(path))
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n".join(parts)


def _extract_docx(path: Path) -> str:
    try:
        import docx
    except ImportError:
        return ""
    document = docx.Document(str(path))
    return "\n".join(p.text for p in document.paragraphs)


def _extract_pptx(path: Path) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        return ""
    presentation = Presentation(str(path))
    parts = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                parts.append(shape.text)
    return "\n".join(parts)
